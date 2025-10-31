# humanitaire.py
# Lit/agrège plusieurs fichiers (appels CSV/XLSX + GRH XLSX), applique un filtre date (si dispo), fusionne un barème de primes, et exporte.

from __future__ import annotations

import os, glob, io, re, unicodedata
from typing import List, Optional

import pandas as pd
import numpy as np

# ──────────────────────────────────────────────────────────────────────────────
# Config par défaut (surcharges via variables d'environnement)
# ──────────────────────────────────────────────────────────────────────────────
DEFAULT_APPELS_GLOB = os.getenv("HUMA_APPELS_GLOB", os.path.join("Data", "*.*"))
DEFAULT_GRH_GLOB    = os.getenv("HUMA_GRH_GLOB",   os.path.join("Data", "extract_grh*.xlsx"))
DEFAULT_BAREME_PATH = os.getenv("HUMA_BAREME_PATH", os.path.join("Data", "Prime don.xlsx"))
DEFAULT_BAREME_SHEET = os.getenv("HUMA_BAREME_SHEET", "Prime dons")
# Taux de conversion € -> DT si la colonne DT n'existe pas dans le barème
EURO_TO_DT = float(os.getenv("HUMA_EURO_TO_DT", "3.32"))

# ──────────────────────────────────────────────────────────────────────────────
# Utils
# ──────────────────────────────────────────────────────────────────────────────
def _normalize_ascii_series(s: pd.Series) -> pd.Series:
    s = s.fillna("").astype(str)
    return s.apply(lambda x: unicodedata.normalize("NFKD", x).encode("ascii", "ignore").decode("ascii"))

def _to_hours(x) -> float:
    """Convertit 'hh:mm[:ss]' ou '2,5' / '2.5' / 2 en heures décimales."""
    if pd.isna(x):
        return 0.0
    s = str(x).strip()
    if re.match(r"^\d{1,2}:\d{2}(:\d{2})?$", s):
        try:
            return pd.to_timedelta(s).total_seconds() / 3600.0
        except Exception:
            return 0.0
    if re.match(r"^\d+,\d+$", s):
        try:
            return float(s.replace(",", "."))
        except Exception:
            return 0.0
    if re.match(r"^\d+(\.\d+)?$", s):
        try:
            return float(s)
        except Exception:
            return 0.0
    return 0.0

def _parse_sheet_date(name: str):
    try:
        return pd.to_datetime(name, format="%Y-%m-%d", errors="coerce")
    except Exception:
        return pd.NaT

def _read_grh_sheet(path: str, sheet_name: str) -> pd.DataFrame:
    try:
        df = pd.read_excel(path, sheet_name=sheet_name, header=6)
    except Exception:
        return pd.DataFrame()

    if df is None or df.empty:
        return pd.DataFrame()

    col_tv  = _pick_col(df, "Agents", "Agent", "TV")
    col_dur = _pick_col(df, "Heure Prod", "Heur Prod", "Durée production", "Duree production", "Heures", "Production")
    if not col_tv or not col_dur:
        return pd.DataFrame()

    out = df[[col_tv, col_dur]].copy()
    out["TV"] = _clean_tv_series(out[col_tv])
    out["Heur Prod"] = out[col_dur].apply(_to_hours)
    out["Heur Prod"] = pd.to_numeric(out["Heur Prod"], errors="coerce").fillna(0.0).astype(float)
    out = out[["TV", "Heur Prod"]]
    out = out[out["TV"] != "INCONNU"]
    return out

def _clean_tv_series(s: pd.Series) -> pd.Series:
    s = s.astype(str).fillna("").str.strip()
    def fix(x: str) -> str:
        xl = x.lower()
        return "INCONNU" if (xl in ("", "nan", "none", "null")) else x
    return s.apply(fix)

def _search_roots() -> List[str]:
    roots = [os.getcwd(), os.path.dirname(__file__), "/data"]
    up = os.path.join(os.getcwd(), "static", "uploads")
    if os.path.isdir(up):
        roots.append(up)
    for dn in ("data", "Data"):
        p = os.path.join(os.getcwd(), dn)
        if os.path.isdir(p):
            roots.append(p)
    return list(dict.fromkeys(roots))

def _resolve_many(glob_pattern: Optional[str], fallback_pattern: str) -> List[str]:
    patterns = [glob_pattern] if glob_pattern else [fallback_pattern]
    roots = _search_roots()
    found: List[str] = []
    for pat in patterns:
        if not pat:
            continue
        if os.path.isabs(pat) or any(ch in pat for ch in (os.sep, "/", "\\")):
            found.extend(glob.glob(pat))
        else:
            for r in roots:
                found.extend(glob.glob(os.path.join(r, pat)))
    found = sorted(set(f for f in found if os.path.exists(f)))
    return found

def _pick_col(df: pd.DataFrame, *cands) -> Optional[str]:
    lower = {c.lower(): c for c in df.columns}
    for n in cands:
        if n.lower() in lower:
            return lower[n.lower()]
    return None

def _detect_date_column(df: pd.DataFrame) -> Optional[str]:
    known = ["DATE", "Date", "Date Appel", "DATE_APPEL", "Jour", "DAY", "DATE APPEL", "CallDate"]
    for name in known:
        col = _pick_col(df, name)
        if col:
            s = _coerce_date_column(df, col)
            if s.notna().mean() >= 0.60:
                return col
    for c in df.columns:
        s = _coerce_date_column(df, c)
        if s.notna().mean() >= 0.60:
            return c
    return None

def _read_any_calls_table(filepath: str) -> pd.DataFrame:
    try:
        ext = os.path.splitext(filepath)[1].lower()
        if ext in [".csv", ".txt"]:
            return pd.read_csv(filepath, sep=";", quotechar='"', encoding="utf-8", dtype=str, engine="python", on_bad_lines="skip")
        else:
            return pd.read_excel(filepath)
    except Exception as e:
        print(f"[HUM] Impossible de lire {os.path.basename(filepath)} : {e}")
        return pd.DataFrame()

# ──────────────────────────────────────────────────────────────────────────────
# Lecture fichiers d'appels (CSV/XLSX), avec filtre date par fichier
# ──────────────────────────────────────────────────────────────────────────────
def _read_all_appels(files: List[str], date_debut=None, date_fin=None) -> pd.DataFrame:
    parts = []
    kept_total = 0
    for f in files:
        base = os.path.basename(f).lower()

        # Ignorer GRH
        if "extract_grh" in base or base.startswith("grh") or re.search(r"\bgrh\b", base):
            _log(f"Ignoré (non-appels): {os.path.basename(f)}")
            continue

        try:
            ext = os.path.splitext(f)[1].lower()
            if ext in (".xlsx", ".xls"):
                df = pd.read_excel(f)
            elif ext == ".csv":
                df = pd.read_csv(f, sep=";", encoding="utf-8", quotechar='"', dtype=str, engine="python", on_bad_lines="skip")
            else:
                _log(f"Ignoré (extension non supportée): {os.path.basename(f)}")
                continue
        except Exception as e:
            print(f"[HUM] Ignoré (appels) {f}: {e}")
            continue

        if df is None or df.empty:
            _log(f"Fichier vide ou illisible: {os.path.basename(f)}")
            continue

        base_code = os.path.basename(f)[:3].upper()
        df["__base_code__"] = base_code

        date_col = _pick_col(df, "DATE") or _pick_col(df, "Date") \
                   or _pick_col(df, "DATE_APPEL", "DATE APPEL", "CallDate") \
                   or _detect_date_column(df)

        _log(f"Fichier: {os.path.basename(f)} | Lignes: {len(df)} | date_col détectée: {date_col}")

        if date_col:
            sdt = _coerce_date_column(df, date_col)
            try:
                min_all = _fmt_dt(sdt.min()); max_all = _fmt_dt(sdt.max())
            except Exception:
                min_all = max_all = "NA"

            _log(f"    Dates détectées (avant filtre): min={min_all}, max={max_all}, NaT%={(sdt.isna().mean()*100):.1f}%")

            if (date_debut or date_fin) and sdt.notna().any():
                mask = pd.Series(True, index=df.index)
                if date_debut:
                    d0 = pd.to_datetime(date_debut); mask &= (sdt >= d0)
                if date_fin:
                    d1 = pd.to_datetime(date_fin) + pd.Timedelta(days=1) - pd.Timedelta(seconds=1)
                    mask &= (sdt <= d1)
                before = len(df)
                df = df[mask]
                after = len(df)
                kept_total += after
                _log(f"    Filtre appliqué [{date_debut} → {date_fin}] : gardées {after}/{before}")
            else:
                kept_total += len(df)
                _log("    Pas de filtre appliqué (pas de dates demandées ou colonne non convertible).")
        else:
            kept_total += len(df)
            _log("    Aucune colonne de date plausible trouvée → pas de filtre (tout est gardé).")

        parts.append(df)

    _log(f"TOTAL lignes appels gardées après concat: {kept_total}")
    return pd.concat(parts, ignore_index=True, sort=False) if parts else pd.DataFrame()

# ──────────────────────────────────────────────────────────────────────────────
# Lecture fichiers GRH (multi-feuilles datées) + agrégation heures
# ──────────────────────────────────────────────────────────────────────────────
def _read_one_grh_any_sheet(path: str) -> pd.DataFrame:
    try:
        df = pd.read_excel(path, sheet_name="Resume", skiprows=6)
        if not df.empty:
            return df
    except Exception:
        pass

    try:
        all_sheets: dict[str, pd.DataFrame] = pd.read_excel(path, sheet_name=None)
        for name, df in all_sheets.items():
            if df is None or df.empty:
                continue
            col_tv  = _pick_col(df, "Agents", "Agent", "TV")
            col_dur = _pick_col(df, "Durée production", "Duree production", "Heur Prod", "Heures", "Production")
            if col_tv and col_dur:
                return df
    except Exception as e:
        print(f"[HUM] Ignoré (GRH lecture multi-feuilles) {path}: {e}")

    return pd.DataFrame()

def _read_all_grh(files: List[str], date_debut: Optional[str] = None, date_fin: Optional[str] = None) -> pd.DataFrame:
    d0 = pd.to_datetime(date_debut).normalize() if date_debut else None
    d1 = pd.to_datetime(date_fin).normalize()   if date_fin   else None

    single_day = None
    if d0 is not None and d1 is not None and d0 == d1:
        single_day = d0

    parts = []
    for f in files:
        try:
            with pd.ExcelFile(f) as xls:
                for sheet in xls.sheet_names:
                    sd = _parse_sheet_date(sheet)
                    if pd.isna(sd):
                        continue
                    sd = sd.normalize()
                    if single_day is not None:
                        if sd != single_day:
                            continue
                    else:
                        if d0 is not None and sd < d0:
                            continue
                        if d1 is not None and sd > d1:
                            continue

                    df_one = _read_grh_sheet(f, sheet)
                    if not df_one.empty:
                        df_one["__sheet_date__"] = sd
                        parts.append(df_one)
        except Exception as e:
            print(f"[HUM] Ignoré (GRH {f}): {e}")

    if not parts:
        _log("GRH: aucune feuille retenue (vérifie les noms de feuilles et les dates filtrées).")
        return pd.DataFrame(columns=["TV", "Heur Prod"])

    df_all = pd.concat(parts, ignore_index=True, sort=False)
    df_all["Heur Prod"] = pd.to_numeric(df_all["Heur Prod"], errors="coerce").fillna(0.0).astype(float)
    df_all["TV"] = _clean_tv_series(df_all["TV"])
    df_all = df_all.groupby("TV", as_index=False, dropna=False)["Heur Prod"].sum()
    return df_all

# ──────────────────────────────────────────────────────────────────────────────
# Barème primes
# ──────────────────────────────────────────────────────────────────────────────
def _charger_bareme_primes(path: str = DEFAULT_BAREME_PATH, sheet: str = DEFAULT_BAREME_SHEET) -> pd.DataFrame:
    """
    Charge le barème de primes.
    Colonnes attendues (après normalisation) :
    - OBJECTIF_DONS
    - DON_MOYEN_CIBLE
    - PRIME_BASE_EUR
    - PRIME_BASE_DT (optionnelle ; si absente, calculée avec EURO_TO_DT)
    """
    try:
        df = pd.read_excel(path, sheet_name=sheet)
    except Exception as e:
        _log(f"AVERTISSEMENT: barème primes introuvable ({path} / {sheet}) : {e}. Primes=0.")
        return pd.DataFrame(columns=["OBJECTIF_DONS","DON_MOYEN_CIBLE","PRIME_BASE_EUR","PRIME_BASE_DT"])

    if df is None or df.empty:
        _log("AVERTISSEMENT: barème vide. Primes=0.")
        return pd.DataFrame(columns=["OBJECTIF_DONS","DON_MOYEN_CIBLE","PRIME_BASE_EUR","PRIME_BASE_DT"])

    # Normalisation des noms
    cols = [c.strip().lower() for c in df.columns]
    df.columns = cols

    # Mapping flexible des colonnes possibles
    col_dons = next((c for c in df.columns if "nombre" in c and "don" in c), None) or next((c for c in df.columns if "dons" in c), None)
    col_dmoy = next((c for c in df.columns if "don" in c and "moyen" in c), None)
    col_pr_eur = next((c for c in df.columns if "prime" in c and "eur" in c), None) or next((c for c in df.columns if "prime" in c and ("euro" in c or "€" in c)), None)
    col_pr_dt  = next((c for c in df.columns if "prime" in c and ("dt" in c or "dinar" in c)), None)

    out = pd.DataFrame()
    try:
        out["OBJECTIF_DONS"]   = pd.to_numeric(df[col_dons], errors="coerce").astype("Int64")
        out["DON_MOYEN_CIBLE"] = pd.to_numeric(df[col_dmoy], errors="coerce").astype("Int64")
        out["PRIME_BASE_EUR"]  = pd.to_numeric(df[col_pr_eur], errors="coerce").fillna(0.0).astype(float)
        if col_pr_dt:
            out["PRIME_BASE_DT"] = pd.to_numeric(df[col_pr_dt], errors="coerce").fillna(0.0).astype(float)
        else:
            out["PRIME_BASE_DT"] = (out["PRIME_BASE_EUR"] * EURO_TO_DT).round(2)
    except Exception as e:
        _log(f"AVERTISSEMENT: colonnes barème non reconnues ({e}). Primes=0.")
        return pd.DataFrame(columns=["OBJECTIF_DONS","DON_MOYEN_CIBLE","PRIME_BASE_EUR","PRIME_BASE_DT"])

    # Nettoyage
    out = out.dropna(subset=["OBJECTIF_DONS","DON_MOYEN_CIBLE"]).copy()
    out["OBJECTIF_DONS"]   = out["OBJECTIF_DONS"].astype(int)
    out["DON_MOYEN_CIBLE"] = out["DON_MOYEN_CIBLE"].astype(int)
    return out

def _nearest(value: float, candidates: List[int]) -> int:
    """Retourne le candidat entier le plus proche de value (avec clipping aux bornes)."""
    if not candidates:
        return int(round(value))
    cands = sorted(set(int(x) for x in candidates))
    lo, hi = cands[0], cands[-1]
    if value <= lo: return lo
    if value >= hi: return hi
    # plus proche
    return min(cands, key=lambda x: abs(x - value))

# ──────────────────────────────────────────────────────────────────────────────
# Cœur
# ──────────────────────────────────────────────────────────────────────────────
def generer_dashboard_humanitaire_df(
    appels_glob: Optional[str] = None,
    grh_glob: Optional[str] = None,
    date_debut: Optional[str] = None,
    date_fin: Optional[str] = None,
    base_code: Optional[str] = None,
) -> pd.DataFrame:
    """
    Agrège *tous* les fichiers d'appels et GRH trouvés.
    Retourne un DF par TV avec colonnes:
      ["ID_TV","TV","Heur Prod","Cu","Don","Don en ligne","Indecis","Montant_Don",
       "Fich_T","Don Moyen","Tx_d’accord","Cu/H","Tx_Argu","OBJECTIF_DONS","DON_MOYEN_CIBLE",
       "PRIME_BASE_EUR","PRIME_BASE_DT"]
    """
    # 1) Localiser
    appels_files = _resolve_many(appels_glob, DEFAULT_APPELS_GLOB)
    grh_files    = _resolve_many(grh_glob,    DEFAULT_GRH_GLOB)

    _log(f"date_debut={date_debut} | date_fin={date_fin}")
    _log(f"APPELS pattern: {appels_glob or DEFAULT_APPELS_GLOB}")
    _log(f"APPELS trouvés: {len(appels_files)} -> {appels_files}")
    _log(f"GRH    pattern: {grh_glob or DEFAULT_GRH_GLOB}")
    _log(f"GRH    trouvés: {len(grh_files)} -> {grh_files}")

    if not appels_files:
        raise FileNotFoundError(f"Aucun fichier d'appels trouvé (pattern: {appels_glob or DEFAULT_APPELS_GLOB})")
    if not grh_files:
        raise FileNotFoundError(f"Aucun fichier GRH trouvé (pattern: {grh_glob or DEFAULT_GRH_GLOB})")

    # 2) Appels
    df_appels = _read_all_appels(appels_files, date_debut, date_fin)
    if base_code:
        base_code = str(base_code).strip().upper()
        if "__base_code__" in df_appels.columns:
            before = len(df_appels)
            df_appels = df_appels[df_appels["__base_code__"] == base_code]
            _log(f"Filtre base_code={base_code}: gardées {len(df_appels)}/{before} lignes d'appels")
        else:
            _log("Avertissement: colonne __base_code__ absente (aucun marquage de base).")

    _log(f"APPELS chargés (après exclusions et filtres): {len(df_appels)} lignes")

    # 3) GRH
    df_grh = _read_all_grh(grh_files, date_debut=date_debut, date_fin=date_fin)
    try:
        _log(f"GRH retenu: {len(df_grh)} TV | Heur Prod total = {df_grh['Heur Prod'].sum():.2f} h")
    except Exception:
        pass

    if df_appels.empty:
        return pd.DataFrame(columns=[
            "ID_TV","TV","Heur Prod","Cu","Don","Don en ligne","Indecis","Montant_Don","Fich_T",
            "Don Moyen","Tx_d’accord","Cu/H","Tx_Argu","OBJECTIF_DONS","DON_MOYEN_CIBLE","PRIME_BASE_EUR","PRIME_BASE_DT"
        ])

    # 4) Colonnes clés
    col_LOG     = _pick_col(df_appels, "LOG", "Log", "Id_tv", "ID_TV")
    col_AGENT   = _pick_col(df_appels, "AGENT", "Agent", "TV")
    col_STATUS  = _pick_col(df_appels, "LIB_STATUS", "Status", "STATUT", "Lib_Status", "LIB_STATUT")
    col_DON_EUR = _pick_col(df_appels, "Don", "Montant_Don", "Montant don", "Montant", "Montant (€)", "MONTANT_DON")

    req = [col_AGENT, col_STATUS, col_DON_EUR]
    if any(c is None for c in req):
        raise KeyError(f"Colonnes requises manquantes dans les fichiers d'appels. Colonnes: {list(df_appels.columns)}")

    # 6) Normalisations
    df_appels["TV_KEY"] = _clean_tv_series(df_appels[col_AGENT])

    montant_raw = (
        df_appels[col_DON_EUR]
        .astype(str)
        .str.replace(r"[^\d,.\-]", "", regex=True)
        .str.replace(",", ".", regex=False)
        .str.strip()
    )
    df_appels["Montant_Don"] = pd.to_numeric(montant_raw, errors="coerce").fillna(0.0)

    col_LIB_DETAIL = _pick_col(df_appels, "LIB_DETAIL", "Lib_Detail", "Detail", "Détail", "LIB DETAIL", "Détail appel")
    col_STATUS     = _pick_col(df_appels, "LIB_STATUS", "Status", "STATUT", "Lib_Status", "LIB_STATUT")

    if col_LIB_DETAIL or col_STATUS:
        s1 = _normalize_ascii_series(df_appels[col_LIB_DETAIL]).str.lower() if col_LIB_DETAIL else pd.Series([""]*len(df_appels))
        s2 = _normalize_ascii_series(df_appels[col_STATUS]).str.lower()     if col_STATUS     else pd.Series([""]*len(df_appels))
        s_all = (s1 + " " + s2).str.strip()

        is_don_mail  = (
            s_all.str.contains(r"\bdon\s*mail\b")
            | s_all.str.contains(r"\bdon\s*par\s*email\b")
            | s_all.str.contains(r"\bdon\s*en\s*ligne\b")
            | s_all.str.contains(r"\blien\s*(de|pour)?\s*don\b")
        )
        is_don_dam   = (
            s_all.str.contains(r"\bdam\b")
            | s_all.str.contains(r"\bdon\s+avec\s+montant\b")
            | (s_all.str.contains(r"\bdon\b") & s_all.str.contains(r"\bmontant\b"))
        )
        is_indecis   = s_all.str.contains(r"\bind[ée]cis\b")
        is_ref_refus = s_all.str.contains(r"\bref\s*refus\b")
        is_cu = is_don_dam | is_don_mail | is_indecis | is_ref_refus

        df_appels["is_Cu"]      = is_cu.astype(int)
        df_appels["is_Don"]     = is_don_dam.astype(int)
        df_appels["is_DonMail"] = is_don_mail.astype(int)
        df_appels["is_Indecis"] = is_indecis.astype(int)
    else:
        s_norm = _normalize_ascii_series(df_appels[col_STATUS]).str.lower()
        col_DETAIL = _pick_col(df_appels, "LIB_DETAIL", "Lib_Detail", "DETAIL")
        d_norm = _normalize_ascii_series(df_appels[col_DETAIL]).str.lower() if col_DETAIL else pd.Series("", index=df_appels.index)

        is_don_mail = (
            s_norm.str.contains(r"\bdon\s*en\s*ligne\b")
            | s_norm.str.contains(r"\bdon\s*par\s*email\b")
            | s_norm.str.contains(r"\blien\s*(de|pour)\s*don\b")
            | s_norm.str.contains(r"\bemail\b")
            | d_norm.str.contains(r"\bdon\s*en\s*ligne\b|\bemail\b|\blien\s*(de|pour)\s*don\b")
        )
        is_don = (
            s_norm.str.contains(r"\bdam\b")
            | s_norm.str.contains(r"\bdon\s*(avec|avc)?\s*montant\b")
            | s_norm.str.contains(r"\biban\b|\bsepa\b|\bprelevement\b|\bvalidation\s*iban\b")
            | d_norm.str.contains(r"\biban\b|\bsepa\b|\bprelevement\b|\bvalidation\s*iban\b")
        ) & (~is_don_mail)
        is_indecis = ((s_norm.str.contains(r"\bindecis\b") & s_norm.str.contains(r"\bdon\b"))
                     | (d_norm.str.contains(r"\bindecis\b") & d_norm.str.contains(r"\bdon\b")))
        is_refus = s_norm.str.contains(r"\brefus") | d_norm.str.contains(r"\brefus")

        df_appels["is_Cu"]      = (is_don | is_don_mail | is_indecis | is_refus).astype(int)
        df_appels["is_Don"]     = is_don.astype(int)
        df_appels["is_DonMail"] = is_don_mail.astype(int)
        df_appels["is_Indecis"] = is_indecis.astype(int)

    # 7) GRH → heures
    if df_grh.empty:
        df_grh_clean = pd.DataFrame(columns=["TV", "Heur Prod"])
    elif set(map(str.lower, df_grh.columns)) >= {"tv", "heur prod"}:
        df_grh_clean = df_grh[["TV", "Heur Prod"]].copy()
        df_grh_clean["TV"] = _clean_tv_series(df_grh_clean["TV"])
        df_grh_clean["Heur Prod"] = pd.to_numeric(df_grh_clean["Heur Prod"], errors="coerce").fillna(0.0).astype(float)
    else:
        col_GRH_TV  = _pick_col(df_grh, "Agents", "Agent", "TV")
        col_GRH_DUR = _pick_col(df_grh, "Durée production", "Duree production", "Heur Prod", "Heures", "Production")
        if col_GRH_TV is None or col_GRH_DUR is None:
            df_grh_clean = pd.DataFrame(columns=["TV", "Heur Prod"])
        else:
            tmp = df_grh[[col_GRH_TV, col_GRH_DUR]].copy()
            tmp["Heur Prod"] = tmp[col_GRH_DUR].apply(_to_hours)
            tmp["TV"] = _clean_tv_series(tmp[col_GRH_TV])
            df_grh_clean = tmp.groupby("TV", as_index=False, dropna=False)["Heur Prod"].sum()

    # 8) Agrégation par TV
    fich_t = df_appels.groupby("TV_KEY", dropna=False).size().rename("Fich_T")

    def _first_non_empty(s):
        s = s.astype(str); s = s[s.str.strip() != ""]
        return s.iloc[0] if len(s) else np.nan

    if col_LOG:
        id_tv = df_appels.groupby("TV_KEY", dropna=False)[col_LOG].apply(_first_non_empty).rename("ID_TV")
    else:
        id_tv = pd.Series(np.nan, index=fich_t.index, name="ID_TV")

    sums = df_appels.groupby("TV_KEY", dropna=False)[
        ["is_Cu", "is_Don", "is_DonMail", "is_Indecis", "Montant_Don"]
    ].sum(numeric_only=True)

    agg = (
        pd.concat([id_tv, sums, fich_t], axis=1)
        .reset_index()
        .rename(columns={
            "TV_KEY": "TV",
            "is_Cu": "Cu",
            "is_Don": "Don",
            "is_DonMail": "Don en ligne",
            "is_Indecis": "Indecis"
        })
    )

    # 9) Join GRH (heures)
    agg = agg.merge(df_grh_clean, on="TV", how="left")
    agg["Heur Prod"] = agg["Heur Prod"].fillna(0.0)

    # Cast sécurité
    for col in ["Cu", "Don", "Don en ligne", "Indecis", "Montant_Don", "Fich_T", "Heur Prod"]:
        if col in agg.columns:
            agg[col] = pd.to_numeric(agg[col], errors="coerce").fillna(0.0).astype(float)

    # 10) KPIs
    denom_dons   = (agg["Don"] + agg["Don en ligne"]).replace(0, np.nan)
    denom_heures = agg["Heur Prod"].replace(0, np.nan)
    denom_fich   = agg["Fich_T"].replace(0, np.nan)
    denom_total  = agg["Cu"].replace(0, np.nan)

    agg["Don Moyen"]   = (agg["Montant_Don"] / denom_dons).fillna(0.0)
    agg["Tx_d’accord"] = ((agg["Don"] + agg["Don en ligne"]) / denom_total).fillna(0.0)
    agg["Cu/H"]        = (agg["Cu"] / denom_heures).fillna(0.0)
    agg["Tx_Argu"]     = (agg["Cu"] / denom_fich).fillna(0.0)

    # ── Barème primes : OBJECTIF_DONS + DON_MOYEN_CIBLE ───────────────────────
    bareme = _charger_bareme_primes()
    if bareme.empty:
        agg["OBJECTIF_DONS"]   = (agg["Don"] + agg["Don en ligne"]).round().astype(int)
        agg["DON_MOYEN_CIBLE"] = agg["Don Moyen"].round().astype(int)
        agg["PRIME_BASE_EUR"]  = 0.0
        agg["PRIME_BASE_DT"]   = 0.0
    else:
        # candidats disponibles dans le barème
        dons_cands = sorted(bareme["OBJECTIF_DONS"].unique().tolist())
        dm_cands   = sorted(bareme["DON_MOYEN_CIBLE"].unique().tolist())

        # objectifs calculés côté prod
        agg["OBJECTIF_DONS"] = (agg["Don"] + agg["Don en ligne"]).round().astype(int)

        # don moyen cible = nearest dans le barème
        agg["_DON_MOY_ROUND_"] = agg["Don Moyen"].round().astype(int)
        agg["DON_MOYEN_CIBLE"] = agg["_DON_MOY_ROUND_"].apply(lambda v: _nearest(v, dm_cands))
        agg.drop(columns=["_DON_MOY_ROUND_"], inplace=True)

        # Clip des dons aux bornes barème pour éviter non-matching
        if dons_cands:
            min_d, max_d = dons_cands[0], dons_cands[-1]
            agg["OBJECTIF_DONS"] = agg["OBJECTIF_DONS"].clip(lower=min_d, upper=max_d)

        agg = agg.merge(
            bareme,
            on=["OBJECTIF_DONS","DON_MOYEN_CIBLE"],
            how="left",
            suffixes=("", "_BAR")
        )

        for k in ["PRIME_BASE_EUR","PRIME_BASE_DT"]:
            if k not in agg.columns:
                agg[k] = 0.0
            agg[k] = pd.to_numeric(agg[k], errors="coerce").fillna(0.0).astype(float)

    # Ordre des colonnes
    cols = [
        "ID_TV","TV","Heur Prod","Cu","Don","Don en ligne","Indecis","Montant_Don","Fich_T",
        "Don Moyen","Tx_d’accord","Cu/H","Tx_Argu",
        "OBJECTIF_DONS","DON_MOYEN_CIBLE","PRIME_BASE_EUR","PRIME_BASE_DT"
    ]
    for c in cols:
        if c not in agg.columns:
            agg[c] = 0.0 if c not in ("ID_TV","TV") else ""
    agg = agg[cols].sort_values(["Don","Don en ligne","Cu"], ascending=False, ignore_index=True)

    # Beautify ID_TV
    try:
        if "ID_TV" in agg.columns and pd.api.types.is_numeric_dtype(agg["ID_TV"]):
            agg["ID_TV"] = agg["ID_TV"].apply(lambda x: int(x) if pd.notna(x) and float(x).is_integer() else x)
    except Exception:
        pass

    # Nettoyage final
    if "TV" in agg.columns:
        before = len(agg)
        agg = agg[~agg["TV"].astype(str).str.strip().isin(["", "nan", "NaN", "INCONNU"])].copy()
        after = len(agg)
        _log(f"Lignes nettoyées (INCONNU/vides) : supprimées {before - after}")

    # Arrondis d'affichage
    if "Heur Prod" in agg.columns:
        agg["Heur Prod"] = agg["Heur Prod"].round(2)
    if "Tx_d’accord" in agg.columns:
        if (agg["Tx_d’accord"].max() <= 1.0):
            agg["Tx_d’accord"] = (agg["Tx_d’accord"] * 100)
        agg["Tx_d’accord"] = agg["Tx_d’accord"].round(1)
    if "Cu/H" in agg.columns:
        agg["Cu/H"] = agg["Cu/H"].round(2)
    if "Don Moyen" in agg.columns:
        agg["Don Moyen"] = agg["Don Moyen"].round(2)
    if "Tx_Argu" in agg.columns:
        agg["Tx_Argu"] = agg["Tx_Argu"].round(2)
    if "PRIME_BASE_EUR" in agg.columns:
        agg["PRIME_BASE_EUR"] = agg["PRIME_BASE_EUR"].round(2)
    if "PRIME_BASE_DT" in agg.columns:
        agg["PRIME_BASE_DT"] = agg["PRIME_BASE_DT"].round(2)

# ─── Nettoyage final anti-NaN avant retour au front ─────────────────────────
# Remplace tous les NaN par None (s'affichera vide côté Jinja)
    agg = agg.where(pd.notna(agg), None)

    # Si tu veux des zéros pour les valeurs numériques (optionnel)
    cols_num = [
        "OBJECTIF_DONS","DON_MOYEN_CIBLE","PRIME_BASE_EUR","PRIME_BASE_DT",
        "Heur Prod","Don Moyen","Tx_d’accord","Cu/H","Tx_Argu",
        "Cu","Don","Don en ligne","Indecis","Montant_Don","Fich_T"
    ]
    for c in cols_num:
        if c in agg.columns:
            agg[c] = pd.to_numeric(agg[c], errors="coerce").fillna(0)


    return agg

# ──────────────────────────────────────────────────────────────────────────────
# Export
# ──────────────────────────────────────────────────────────────────────────────
def export_dashboard_humanitaire_xlsx(
    appels_glob: Optional[str] = None,
    grh_glob: Optional[str] = None,
    date_debut: Optional[str] = None,
    date_fin: Optional[str] = None,
    base_code: Optional[str] = None,
) -> io.BytesIO:
    df = generer_dashboard_humanitaire_df(appels_glob, grh_glob, date_debut, date_fin, base_code=base_code)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Dashboard_Humanitaire")
    output.seek(0)
    return output

# ──────────────────────────────────────────────────────────────────────────────
# Utilitaires divers
# ──────────────────────────────────────────────────────────────────────────────
def extract_tv_list(appels_glob: Optional[str] = None) -> List[str]:
    files = _resolve_many(appels_glob, DEFAULT_APPELS_GLOB)
    if not files:
        return []
    parts = []
    for f in files:
        try:
            base = os.path.basename(f).lower()
            if base.endswith(".csv"):
                parts.append(pd.read_csv(f, sep=";", encoding="utf-8", quotechar='"', dtype=str, engine="python", on_bad_lines="skip"))
            elif base.endswith(".xlsx") or base.endswith(".xls"):
                parts.append(pd.read_excel(f))
        except Exception:
            pass
    if not parts:
        return []
    df = pd.concat(parts, ignore_index=True, sort=False)
    col_AGENT = _pick_col(df, "AGENT", "Agent", "TV")
    if not col_AGENT:
        return []
    tv = _clean_tv_series(df[col_AGENT])
    tv = tv[(~tv.str.upper().eq("INCONNU")) & (tv.str.strip() != "")]
    return sorted(tv.unique().tolist())

def _excel_or_text_to_datetime(s: pd.Series) -> pd.Series:
    s_str = s.astype(str).str.strip()
    s_txt = pd.to_datetime(s_str, errors="coerce", dayfirst=True, infer_datetime_format=True)
    s_num = pd.to_numeric(s_str, errors="coerce")
    mask_yyyymmdd = s_str.str.match(r"^(19|20)\d{6}$")
    s_ymd = pd.to_datetime(s_str.where(mask_yyyymmdd), format="%Y%m%d", errors="coerce")
    s_txt = s_ymd.combine_first(s_txt)
    plausible_excel = s_num.between(30000, 80000)
    if plausible_excel.mean() > 0.4:
        s_excel = pd.to_datetime(s_num, unit="D", origin="1899-12-30", errors="coerce")
        try:
            many_1970 = s_txt.notna().mean() > 0 and (s_txt.dt.year == 1970).mean() > 0.6
        except Exception:
            many_1970 = False
        if s_txt.isna().mean() > 0.4 or many_1970:
            s_txt = s_excel.where(plausible_excel, s_txt).combine_first(s_excel)
    return s_txt

def _coerce_date_column(df: pd.DataFrame, colname: str) -> pd.Series:
    return _excel_or_text_to_datetime(df[colname])

def _fmt_dt(x):
    try:
        return pd.to_datetime(x).strftime("%Y-%m-%d")
    except Exception:
        return str(x)

def _log(msg: str):
    print(f"[HUM/DEBUG] {msg}")

def list_bases_disponibles(appels_glob: Optional[str] = None) -> List[str]:
    patterns: List[str] = []
    if appels_glob:
        patterns = [appels_glob]
    else:
        patterns = [
            os.path.join("Data", "*.xlsx"),
            os.path.join("Data", "*.xls"),
            os.path.join("Data", "*.csv"),
        ]
    files: List[str] = []
    for pat in patterns:
        files.extend(_resolve_many(pat, pat))
    bases = []
    for f in files:
        nm = os.path.basename(f)
        low = nm.lower()
        if "grh" in low or low.startswith("extract_grh"):
            continue
        if len(nm) >= 3:
            bases.append(nm[:3].upper())
    return sorted(set(bases))

def export_dashboard_humanitaire_xlsx_from_df(df: pd.DataFrame, out_path: str | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "Tableau"

    headers = list(df.columns)
    ws.append(headers)

    for _, row in df.iterrows():
        ws.append([row.get(h, "") for h in headers])

    for idx, col in enumerate(headers, start=1):
        max_len = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).tolist()]) if col in df else len(str(col))
        ws.column_dimensions[get_column_letter(idx)].width = min(60, max(10, max_len + 2))

    bio = BytesIO()
    wb.save(bio)
    data = bio.getvalue()
    if out_path:
        with open(out_path, "wb") as f:
            f.write(data)
    return data

def extract_tv_list_from_df(df: pd.DataFrame) -> list[str]:
    if "TV" not in df.columns:
        return []
    tvs = df["TV"].astype(str).fillna("").str.strip()
    tvs = [t for t in tvs.unique().tolist() if t]
    return sorted(tvs, key=lambda s: s.lower())
